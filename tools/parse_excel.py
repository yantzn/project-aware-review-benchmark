from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def parse_workbook(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=False)
    sheets: list[dict[str, object]] = []

    for worksheet in workbook.worksheets:
        cells: list[dict[str, object]] = []
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cells.append(
                    {
                        "coordinate": cell.coordinate,
                        "value": cell.value,
                        "data_type": cell.data_type,
                    }
                )

        hidden_rows = [
            index for index, dimension in worksheet.row_dimensions.items() if dimension.hidden
        ]
        hidden_columns = [
            key for key, dimension in worksheet.column_dimensions.items() if dimension.hidden
        ]
        sheets.append(
            {
                "sheet": worksheet.title,
                "state": worksheet.sheet_state,
                "merged_ranges": [str(item) for item in worksheet.merged_cells.ranges],
                "hidden_rows": hidden_rows,
                "hidden_columns": hidden_columns,
                "cells": cells,
            }
        )

    return {
        "source_file": path.as_posix(),
        "parser": "openpyxl",
        "limitations": [
            "図形・画像の意味は構造化しない",
            "色・罫線など視覚表現の意味は判定しない",
            "数式は式文字列を保持し、計算結果の正しさは判定しない",
        ],
        "sheets": sheets,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    print(json.dumps(parse_workbook(args.workbook), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
